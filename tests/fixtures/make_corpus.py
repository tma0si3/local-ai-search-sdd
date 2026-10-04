"""Generates the fixture corpus (T004). Run: uv run python tests/fixtures/make_corpus.py

Committed as a script rather than binary blobs so the fixtures are inspectable and
regenerable — a reviewer can see exactly what the tests are searching.
"""

from __future__ import annotations

from pathlib import Path

CORPUS = Path(__file__).parent / "corpus"

SOIL = """Field Study 2025: Soil Erosion on the North Ridge

Across all three monitored plots, erosion rates accelerated sharply following the removal of
hedgerow cover in early spring. Plot B lost an estimated 14 tonnes of topsoil per hectare over
the season, roughly triple the ten-year average for comparable terrain.

The dominant mechanism was rill formation during high-intensity rainfall. Once a rill network
established itself, subsequent moderate rainfall events caused disproportionate losses.

We conclude that hedgerow retention is the single most effective intervention available at this
site, and that replanting should be prioritised on the steepest gradients before the next wet
season.
"""

BUDGET = """Quarterly Budget Review

Operating expenditure came in 8 percent under forecast this quarter, largely because two planned
equipment purchases were deferred.

Staff costs rose modestly following the annual review. Travel remained well below pre-pandemic
levels and we see no sign of a return to earlier patterns.

Recommendation: carry the underspend forward rather than redistributing it, since the deferred
equipment will fall due next quarter.
"""

RECIPE = """# Sourdough Notes

## Starter

Feed twice daily at room temperature. A healthy starter doubles in four to six hours and smells
sharp rather than sour.

## Hydration

Seventy-five percent hydration gives an open crumb but a slack dough that is harder to shape.
Beginners should start nearer sixty-five percent.

## Baking

Bake covered for twenty minutes, then uncovered for a further twenty-five. The crust should be
deep brown, almost approaching burnt at the edges.
"""

MEETING = """Project Kickoff Meeting

Attendees agreed the first milestone is a working prototype by the end of the month.

Open questions concerned data retention and whether the archive needs to be searchable from day
one. No decision was reached and this was carried over.

Action: circulate a short written proposal before the next meeting.
"""


def write_pdf(path: Path, title: str, body: str) -> None:
    """Minimal single-page PDF written by hand — avoids a generation dependency."""
    text = f"{title}\n\n{body}"
    lines = [line for line in text.splitlines() if line.strip()]

    content_lines = ["BT", "/F1 11 Tf", "50 760 Td", "14 TL"]
    for line in lines:
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content_lines.append(f"({escaped}) Tj T*")
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("latin-1", errors="replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body_bytes in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body_bytes + b"\nendobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode()

    path.write_bytes(bytes(out))


def write_docx(path: Path, title: str, body: str) -> None:
    import docx

    document = docx.Document()
    document.add_heading(title, level=1)
    for paragraph in body.split("\n\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    document.save(str(path))


def main() -> None:
    CORPUS.mkdir(parents=True, exist_ok=True)
    (CORPUS / "nested" / "deeper").mkdir(parents=True, exist_ok=True)

    write_pdf(CORPUS / "field-study-2025.pdf", "Field Study 2025", SOIL)
    (CORPUS / "budget-review.txt").write_text(BUDGET, encoding="utf-8")
    (CORPUS / "sourdough-notes.md").write_text(RECIPE, encoding="utf-8")
    write_docx(CORPUS / "nested" / "deeper" / "kickoff-meeting.docx", "Project Kickoff", MEETING)

    # Edge cases: each must be reported as skipped with a reason, never silently dropped.
    (CORPUS / "empty.txt").write_text("", encoding="utf-8")
    (CORPUS / "whitespace-only.txt").write_text("\n\n   \t\n", encoding="utf-8")
    (CORPUS / "notes.pages").write_text("unsupported format", encoding="utf-8")

    print(f"Wrote fixture corpus to {CORPUS}")


if __name__ == "__main__":
    main()

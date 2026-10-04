"""Reading text out of the four supported document formats.

No I/O beyond reading the file it is handed, and no knowledge of the index. Failures raise
``ExtractionError`` carrying a reason written for the user, because that reason is shown
verbatim in the skipped-documents list (FR-010, SC-007).

Markdown is indexed as raw text. Its syntax markers are sparse enough not to disturb
embeddings meaningfully, and skipping a parser keeps a dependency out (research.md §3).
"""

from __future__ import annotations

from pathlib import Path

from .errors import ExtractionError

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".txt": "txt", ".md": "md", ".docx": "docx"}


def file_type_of(path: Path) -> str | None:
    """Return the document type, or None if this file is not one we handle."""
    return SUPPORTED_EXTENSIONS.get(path.suffix.lower())


def extract_text(path: Path) -> str:
    """Return the document's text.

    Raises:
        ExtractionError: with a reason suitable for display to the user.
    """
    file_type = file_type_of(path)
    if file_type is None:
        raise ExtractionError(
            f"'{path.suffix or 'no extension'}' files are not supported. "
            "Supported formats are PDF, TXT, Markdown, and Word (.docx)."
        )

    extractors = {
        "pdf": _extract_pdf,
        "docx": _extract_docx,
        "txt": _extract_plain_text,
        "md": _extract_plain_text,
    }

    try:
        text = extractors[file_type](path)
    except ExtractionError:
        raise
    except PermissionError as exc:
        raise ExtractionError(
            "This file could not be read because permission was denied. Check its "
            "permissions and re-index."
        ) from exc
    except Exception as exc:
        raise ExtractionError(
            f"This file could not be read: {exc}. It may be corrupt or in an unexpected format."
        ) from exc

    if not text.strip():
        raise ExtractionError(
            "No readable text was found in this file. If it is a scanned document, it "
            "contains images rather than text; this application does not perform OCR."
        )
    return text


def _extract_plain_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        # Fall back rather than skip. A mis-decoded character is a far smaller loss than
        # dropping the whole document.
        return path.read_text(encoding="utf-8", errors="replace")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ExtractionError(
            "This PDF is password-protected and cannot be read. Remove the password and "
            "re-index if you want it included."
        )
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(page for page in pages if page.strip())


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n\n".join(block for block in blocks if block.strip())
